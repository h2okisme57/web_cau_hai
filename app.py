import io
import time
import zipfile
import re
import xml.etree.ElementTree as ET
import streamlit as st
from google import genai
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt

# Danh sách model theo thứ tự ưu tiên
MODEL_CANDIDATES = [
    "gemini-2.0-flash",
    "gemini-2.0-flash-exp",
    "gemini-1.5-flash",
    "gemini-3.8-flash",
]

st.set_page_config(
    page_title="AI Giáo Dục - Lê Minh Tuấn",
    page_icon="🎓",
    layout="wide"
)

# Header định danh tác giả
st.markdown("## 🎓 HỆ THỐNG TRỢ LÝ AI GIÁO DỤC TOÀN DIỆN")
st.markdown("**Tác giả:** Lê Minh Tuấn - GV (Chuyên môn: Lịch sử & Địa lí)")
st.caption("Chương trình chuyển giao kỹ thuật ứng dụng AI trong GD – Tỉnh Vĩnh Long (10/2026)")
st.divider()

# Sidebar: Thiết lập API & Thông tin bài học
with st.sidebar:
    st.header("⚙️ Thiết lập hệ thống")
    api_key_input = st.text_input("Nhập Google Gemini API Key:", type="password")
    api_key = api_key_input.strip() or st.secrets.get("GEMINI_API_KEY", "").strip()
    
    # Nút kiểm tra API Key tự động dò model đang mở cho tài khoản
    if st.button("🔍 Kiểm tra API Key", use_container_width=True):
        if not api_key:
            st.error("Chưa nhập API Key!")
        else:
            with st.spinner("Đang kiểm tra kết nối với hệ thống Google AI..."):
                try:
                    client = genai.Client(api_key=api_key)
                    # Quét danh sách model được phép của tài khoản
                    valid_models = [m.name for m in client.models.list()]
                    chosen = None
                    for candidate in MODEL_CANDIDATES:
                        for m_name in valid_models:
                            if candidate in m_name:
                                chosen = m_name
                                break
                        if chosen:
                            break
                    target = chosen if chosen else (valid_models[0] if valid_models else "gemini-2.0-flash")
                    st.session_state["active_model"] = target
                    st.success(f"✅ Kết nối thành công!\n\nModel sẵn sàng: `{target}`")
                except Exception as e:
                    st.error(f"❌ Lỗi xác thực: {e}")

    st.divider()
    st.header("📋 Thông tin bài dạy thực nghiệm")
    mon_hoc = st.selectbox("Môn học:", [
        "Lịch sử và Địa lí (Lịch sử)", "Lịch sử", "Địa lí", "Toán học", 
        "Ngữ văn", "Khoa học tự nhiên", "Tin học", "Giáo dục công dân"
    ], index=0)
    lop = st.selectbox("Khối lớp:", [f"Lớp {i}" for i in range(6, 13)], index=1)
    ten_bai = st.text_input("Tên bài học:", value="Bài 2. Các cuộc phát kiến địa lí")
    so_tiet = st.number_input("Thời lượng (tiết):", min_value=1, max_value=6, value=2)
    yeu_cau_can_dat = st.text_area(
        "Yêu cầu cần đạt (theo CT GDPT 2018):", 
        height=100, 
        value="- Trình bày được nguyên nhân và điều kiện của các cuộc phát kiến địa lí.\n- Mô tả được các cuộc phát kiến địa lí của B. Đi-a-xơ, C. Cô-lôm-bô, V. Ga-ma và Ph. Ma-gien-lan.\n- Đánh giá được tác động của các cuộc phát kiến địa lí đối với tiến trình lịch sử."
    )

# Hàm đọc file docx có Fallback mở trực tiếp gói ZIP XML
def read_uploaded_file(uploaded_file):
    try:
        raw_bytes = uploaded_file.read()
        if uploaded_file.name.endswith(".docx"):
            # Cách 1: Dùng python-docx tiêu chuẩn
            try:
                doc = Document(io.BytesIO(raw_bytes))
                full_text = []
                for p in doc.paragraphs:
                    if p.text.strip():
                        full_text.append(p.text.strip())
                for table in doc.tables:
                    for row in table.rows:
                        row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_text:
                            full_text.append(" | ".join(row_text))
                extracted = "\n".join(full_text)
                if extracted.strip():
                    return extracted
            except Exception:
                pass

            # Cách 2: Fallback mở thẳng package zip xml nếu content-type bị lỗi
            try:
                with zipfile.ZipFile(io.BytesIO(raw_bytes)) as z:
                    # Tìm file văn bản chính của Word
                    target_xml = None
                    for name in z.namelist():
                        if name.endswith("word/document.xml") or name.endswith("document.xml"):
                            target_xml = name
                            break
                    if target_xml:
                        xml_content = z.read(target_xml)
                        tree = ET.fromstring(xml_content)
                        # Trích xuất toàn bộ text trong thẻ w:t
                        texts = [elem.text for elem in tree.iter() if elem.tag.endswith('t') and elem.text]
                        return "\n".join(texts)
            except Exception:
                pass

        elif uploaded_file.name.endswith(".txt"):
            return raw_bytes.decode("utf-8", errors="ignore")

    except Exception as e:
        st.error(f"Lỗi khi đọc file tài liệu: {e}")
        return ""
    return ""

# Hàm gọi Gemini AI tự động thích ứng với model đã nhận diện
def call_gemini(prompt_text, key):
    if not key:
        st.error("Vui lòng nhập API Key ở thanh bên trái!")
        return None
    try:
        client = genai.Client(api_key=key.strip())
        
        # Lấy model đang active hoặc mặc định gemini-2.0-flash
        target_model = st.session_state.get("active_model", "gemini-2.0-flash")
        
        for retry in range(2):
            try:
                response = client.models.generate_content(
                    model=target_model,
                    contents=prompt_text,
                )
                
                # Cách 1: Lấy trực tiếp qua response.text
                if response and hasattr(response, "text") and response.text:
                    return response.text
                
                # Cách 2: Fallback bóc tách từ candidates nếu response.text bị rỗng
                if response and response.candidates:
                    candidate = response.candidates[0]
                    if candidate.content and candidate.content.parts:
                        parts_text = "".join([p.text for p in candidate.content.parts if hasattr(p, "text") and p.text])
                        if parts_text.strip():
                            return parts_text
                    
                    # Nếu bị chặn bởi bộ lọc an toàn của Google
                    finish_reason = getattr(candidate, "finish_reason", None)
                    if finish_reason:
                        st.warning(f"AI ngắt phản hồi do lý do kỹ thuật: {finish_reason}")
                        return None

            except Exception as err:
                err_msg = str(err).lower()
                if "503" in err_msg or "high demand" in err_msg or "unavailable" in err_msg:
                    time.sleep(2)
                    continue
                else:
                    st.error(f"Lỗi API: {err}")
                    return None

        st.warning("AI không trả về nội dung. Vui lòng bấm thử lại hoặc rút gọn bớt nội dung nạp vào.")
        return None
    except Exception as e:
        st.error(f"Lỗi kết nối client AI: {e}")
        return None

# Xuất file Word (.docx)
def export_docx(title, content, author="Lê Minh Tuấn - GV"):
    doc = Document()
    doc.add_heading(title, level=0)
    p = doc.add_paragraph()
    p.add_run(f"Người thực hiện: {author}\n").bold = True
    p.add_run("Quy chuẩn chuyên môn: Công văn 5512/BGDĐT & Công văn 7991/BGDĐT-GDTrH\n").italic = True
    for para in content.split("\n"):
        doc.add_paragraph(para)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

# Hàm xuất PowerPoint cho Bài 3 (TO - VỪA - NHỎ + Speaker Notes)
def export_advanced_pptx(title, raw_text, author):
    prs = Presentation()
    prs.slide_width = Inches(13.333) # Chuẩn màn hình 16:9
    prs.slide_height = Inches(7.5)
    
    # 1. Slide bìa
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title.replace("**", "").replace("#", "").strip()
    slide.placeholders[1].text = f"Môn học: {mon_hoc} - {lop}\nGiáo viên biên soạn: {author}"
    
    # 2. Tách slide linh hoạt bằng Regex (bắt được cả '### SLIDE 1:', '**SLIDE 2**', 'SLIDE:')
    slide_blocks = re.split(r'(?:#{1,4}\s*)?(?:\*\*)?SLIDE\s*\d*[:\-–]?', raw_text, flags=re.IGNORECASE)
    
    for block in slide_blocks:
        block = block.strip()
        if not block:
            continue
            
        # Bỏ qua đoạn giới thiệu mở đầu của AI nếu không chứa nội dung slide
        if "dưới đây là kịch bản" in block.lower() or "powerpoint" in block.lower() and len(block) < 150:
            if not any(k in block.upper() for k in ["TO:", "VỪA:", "NHỎ:"]):
                continue

        lines = [l.strip() for l in block.split("\n") if l.strip()]
        if not lines:
            continue
            
        # Dòng đầu tiên là tiêu đề slide (làm sạch ký tự markdown)
        slide_title = re.sub(r'[*#_]', '', lines[0]).strip()
        body_lines = []
        speaker_notes = []
        is_notes = False
        
        for l in lines[1:]:
            clean_l = re.sub(r'[*#_]', '', l).strip()
            if not clean_l:
                continue
            if any(k in clean_l.upper() for k in ["LỜI DẪN GV:", "GHI CHÚ GV:", "NOTES:"]):
                is_notes = True
                clean_l = re.sub(r'^(LỜI DẪN GV|GHI CHÚ GV|NOTES)[:\-–]\s*', '', clean_l, flags=re.IGNORECASE)
                if clean_l:
                    speaker_notes.append(clean_l)
                continue
                
            if is_notes:
                speaker_notes.append(clean_l)
            else:
                body_lines.append(clean_l)
                
        # Khởi tạo Slide nội dung chuẩn (Title and Content layout)
        new_slide = prs.slides.add_slide(prs.slide_layouts[1])
        new_slide.shapes.title.text = slide_title
        
        tf = new_slide.placeholders[1].text_frame
        tf.clear()
        
        for bl in body_lines:
            bl_clean = bl.lstrip("-•* ")
            if not bl_clean:
                continue
            p = tf.add_paragraph()
            
            # Phân cấp cỡ chữ theo tiêu chuẩn TO - VỪA - NHỎ
            if bl_clean.upper().startswith("TO:"):
                p.text = re.sub(r'^TO[:\-–]\s*', '', bl_clean, flags=re.IGNORECASE)
                p.font.size = Pt(24)
                p.font.bold = True
            elif bl_clean.upper().startswith("VỪA:"):
                p.text = re.sub(r'^VỪA[:\-–]\s*', '', bl_clean, flags=re.IGNORECASE)
                p.font.size = Pt(19)
            elif bl_clean.upper().startswith("NHỎ:"):
                p.text = re.sub(r'^NHỎ[:\-–]\s*', '', bl_clean, flags=re.IGNORECASE)
                p.font.size = Pt(15)
                p.font.italic = True
            else:
                p.text = bl_clean
                p.font.size = Pt(19)
                
        # Đổ lời dẫn giáo viên vào Speaker Notes của slide
        if speaker_notes:
            notes_slide = new_slide.notes_slide
            text_frame = notes_slide.notes_text_frame
            text_frame.text = "\n".join(speaker_notes)
            
    bio = io.BytesIO()
    prs.save(bio)
    return bio.getvalue()
# Giao diện 4 Tabs bài tập
tab1, tab2, tab3, tab4 = st.tabs([
    "📄 1. Kế hoạch bài dạy (Bài 1)", 
    "📝 2. Đề kiểm tra CV 7991 (Bài 1)",
    "💡 3. Trợ lý Sáng kiến kinh nghiệm (Bài 2)",
    "🖥️ 4. Nạp KHBD xuất Slide PPTX (Bài 3)"
])

# TAB 1: KẾ HOẠCH BÀI DẠY (CV 5512)
with tab1:
    st.subheader("Soạn Kế hoạch bài dạy (Giáo án CV 5512)")
    if st.button("🚀 Khởi tạo Kế hoạch bài dạy", key="btn_khbd"):
        with st.spinner("Đang biên soạn KHBD theo đúng chuẩn khung văn bản Bộ GD&ĐT..."):
            prompt_khbd = f"""
            Đóng vai trò là chuyên gia sư phạm. Hãy biên soạn Kế hoạch bài dạy chuẩn cho:
            - Môn: {mon_hoc} - {lop}
            - Tên bài: {ten_bai} (Thời lượng: {so_tiet} tiết)
            - Yêu cầu cần đạt: {yeu_cau_can_dat}
            
            Khung kế hoạch bài dạy chuẩn:
            I. MỤC TIÊU (Năng lực đặc thù, Năng lực chung, Phẩm chất)
            II. THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU
            III. TIẾN TRÌNH DẠY HỌC (Gồm 4 hoạt động: Khởi động, Hình thành kiến thức, Luyện tập, Vận dụng; mỗi hoạt động tổ chức đủ 4 bước: Giao nhiệm vụ -> Thực hiện -> Báo cáo -> Nhận định).
            Ký tên người soạn: Lê Minh Tuấn - GV.
            """
            st.session_state["res_khbd"] = call_gemini(prompt_khbd, api_key)
            
    if st.session_state.get("res_khbd"):
        st.markdown(st.session_state["res_khbd"])
        docx_bytes = export_docx(f"KHBD_{ten_bai}", st.session_state["res_khbd"], "Lê Minh Tuấn - GV")
        st.download_button(
            "📥 Tải về KHBD Word (.docx)", 
            data=docx_bytes, 
            file_name=f"KHBD_LeMinhTuan.docx", 
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

# TAB 2: ĐỀ KIỂM TRA ĐỊNH KỲ THEO CÔNG VĂN 7991
with tab2:
    st.subheader("Tạo Ma trận, Bản đặc tả và Đề kiểm tra chuẩn CV 7991")
    col1, col2 = st.columns(2)
    with col1:
        thoi_gian = st.selectbox("Thời lượng kiểm tra:", ["45 phút (Định kì)", "60 phút", "90 phút"])
    with col2:
        co_tra_loi_ngan = st.checkbox("Môn có trắc nghiệm Trả lời ngắn", value=False)

    if st.button("🚀 Khởi tạo Ma trận & Đề kiểm tra 7991", key="btn_exam_7991"):
        with st.spinner("Đang thiết lập ma trận và đặc tả theo Công văn 7991/BGDĐT-GDTrH..."):
            prompt_7991 = f"""
            Thiết kế bộ hồ sơ kiểm tra định kì chuẩn xác theo CÔNG VĂN SỐ 7991/BGDĐT-GDTrH:
            - Môn: {mon_hoc} - {lop} | Bài: {ten_bai} | Thời gian: {thoi_gian}
            - Giáo viên ra đề: Lê Minh Tuấn - GV
            - Tỉ lệ điểm: Nhiều lựa chọn 30%, Đúng - Sai 20%, {"Trả lời ngắn 20%" if co_tra_loi_ngan else "Chuyển trả lời ngắn sang Đúng-Sai 40%"}, Tự luận 30%.
            - Tỉ lệ nhận thức: Biết 40%, Hiểu 30%, Vận dụng 30%.
            Xuất đầy đủ 4 phần:
            I. KHUNG MA TRẬN ĐỀ KIỂM TRA ĐỊNH KÌ (Bảng Markdown).
            II. BẢN ĐẶC TẢ ĐỀ KIỂM TRA ĐỊNH KÌ (Có Yêu cầu cần đạt và Năng lực).
            III. ĐỀ KIỂM TRA HOÀN CHỈNH.
            IV. ĐÁP ÁN VÀ HƯỚNG DẪN CHẤM (Barem điểm Đúng - Sai chuẩn: đúng 1 ý=0.1đ, 2 ý=0.25đ, 3 ý=0.5đ, 4 ý=1.0đ; barem tự luận chi tiết).
            """
            st.session_state["res_exam"] = call_gemini(prompt_7991, api_key)
            
    if st.session_state.get("res_exam"):
        st.markdown(st.session_state["res_exam"])
        docx_exam_bytes = export_docx(f"De_Kiem_Tra_7991_{ten_bai}", st.session_state["res_exam"], "Lê Minh Tuấn - GV")
        st.download_button(
            "📥 Tải về Bộ Đề kiểm tra chuẩn CV 7991 (.docx)", 
            data=docx_exam_bytes, 
            file_name=f"De_Kiem_Tra_7991_LeMinhTuan.docx", 
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

# TAB 3: TRỢ LÝ SÁNG KIẾN KINH NGHIỆM (BÀI TẬP 2)
with tab3:
    st.subheader("💡 Trợ lý Xây dựng Sáng kiến kinh nghiệm (SKKN)")
    ten_skkn = st.text_input("Tên đề tài SKKN:", value="Ứng dụng Trí tuệ nhân tạo (AI) và Bản đồ số hóa trong đổi mới dạy học phân môn Lịch sử cấp THCS")
    thuc_trang = st.text_area("Thực trạng trước khi áp dụng:", value="Học sinh học lịch sử thụ động, việc ghi nhớ các mốc niên đại và hải trình thám hiểm còn máy móc, trừu tượng.")
    if st.button("🚀 Khởi tạo Dự thảo SKKN hoàn chỉnh", key="btn_skkn"):
        with st.spinner("Đang xây dựng dự thảo SKKN theo thể thức chuẩn Thông tư 18/2013/TT-BKHCN..."):
            prompt_skkn = f"""
            Chấp bút một bản DỰ THẢO SÁNG KIẾN KINH NGHIỆM hoàn chỉnh theo quy định mẫu của Bộ GDĐT:
            - Tên sáng kiến: {ten_skkn}
            - Tác giả: Lê Minh Tuấn - GV | Môn: {mon_hoc}
            - Thực trạng: {thuc_trang}
            Gồm đủ: Đơn yêu cầu công nhận sáng kiến (Phụ lục I Thông tư 18), Lý do chọn đề tài, Mục đích, Giải pháp thực hiện chi tiết, Hiệu quả thu được (bảng đối chứng số liệu) và Bài học kinh nghiệm.
            """
            st.session_state["res_skkn"] = call_gemini(prompt_skkn, api_key)
            
    if st.session_state.get("res_skkn"):
        st.markdown(st.session_state["res_skkn"])
        docx_skkn = export_docx(f"SKKN_{ten_skkn[:30]}", st.session_state["res_skkn"], "Lê Minh Tuấn - GV")
        st.download_button(
            "📥 Tải về Bản Dự thảo SKKN (.docx)", 
            data=docx_skkn, 
            file_name="Du_Thao_SKKN_LeMinhTuan.docx", 
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

# TAB 4: CHUYỂN KHBD THÀNH SLIDE POWERPOINT (BÀI TẬP 3 CHUYÊN SÂU - CÓ NẠP FILE)
with tab4:
    st.subheader("🖥️ Trợ lý Nạp Kế hoạch bài dạy & Chuyển đổi thành Bộ Slide (.PPTX)")
    st.caption("Tiếp nhận tệp KHBD (.docx/.txt), phân tích cấu trúc và thiết kế slide phân cấp TO - VỪA - NHỎ kèm Lời dẫn giáo viên")
    
    # 3 phương thức cung cấp KHBD
    source_mode = st.radio(
        "Chọn phương thức nạp Kế hoạch bài dạy:", 
        ["📁 Tải lên tệp KHBD (.docx, .txt)", "✏️ Dán nội dung văn bản KHBD", "🔄 Sử dụng KHBD vừa tạo ở Tab 1"], 
        horizontal=True
    )
    
    final_khbd_content = ""
    
    if source_mode == "📁 Tải lên tệp KHBD (.docx, .txt)":
        uploaded_doc = st.file_uploader("Chọn file Kế hoạch bài dạy từ máy tính của thầy/cô:", type=["docx", "txt"], key="uploader_khbd")
        if uploaded_doc is not None:
            with st.spinner("Đang trích xuất nội dung văn bản từ tệp..."):
                final_khbd_content = read_uploaded_file(uploaded_doc)
            if final_khbd_content and final_khbd_content.strip():
                st.success(f"✅ Đã nạp thành công tệp: **{uploaded_doc.name}** ({len(final_khbd_content)} ký tự)")
                with st.expander("👁️ Xem trước nội dung đã trích xuất từ tệp"):
                    st.text_area("Nội dung file:", value=final_khbd_content, height=180, disabled=True)
            else:
                st.warning("⚠️ Không thể trích xuất văn bản từ tệp Word này (file có thể bị lỗi định dạng). Vui lòng chuyển sang tab '✏️ Dán nội dung văn bản KHBD' để dán trực tiếp!")
                
    elif source_mode == "✏️ Dán nội dung văn bản KHBD":
        final_khbd_content = st.text_area(
            "Dán toàn bộ nội dung giáo án / Kế hoạch bài dạy vào đây:",
            height=220,
            value=f"KẾ HOẠCH BÀI DẠY: {ten_bai}\nMôn học: {mon_hoc} - {lop}\nThời lượng: {so_tiet} tiết\nYêu cầu cần đạt: {yeu_cau_can_dat}"
        )
        
    elif source_mode == "🔄 Sử dụng KHBD vừa tạo ở Tab 1":
        final_khbd_content = st.session_state.get("res_khbd", "")
        if final_khbd_content:
            st.success("✅ Đã lấy thành công Kế hoạch bài dạy từ Tab 1.")
            with st.expander("👁️ Xem lại nội dung KHBD từ Tab 1"):
                st.markdown(final_khbd_content)
        else:
            st.info("💡 Bạn chưa bấm tạo ở Tab 1. Hãy bấm sang mục Dán văn bản hoặc Tải file lên.")

    col_s1, col_s2 = st.columns(2)
    with col_s1:
        so_slide_target = st.slider("Số lượng Slide mục tiêu:", 6, 16, 8)
    with col_s2:
        kieu_thiet_ke = st.selectbox("Phong cách sư phạm:", ["Chuẩn hóa bám sát tiến trình KHBD", "Tương tác phát vấn & Hoạt động nhóm", "Trực quan hóa trọng tâm"])

    if st.button("🚀 Bắt đầu chuyển đổi KHBD thành Bộ Slide PPTX", type="primary", key="btn_run_bt3"):
        if not final_khbd_content or not final_khbd_content.strip():
            st.error("⚠️ Chưa có nội dung Kế hoạch bài dạy! Vui lòng tải file hoặc chọn mục 'Dán nội dung' để tiếp tục.")
        else:
            with st.spinner("AI đang phân tích tiến trình bài dạy và thiết kế kịch bản trình chiếu phân cấp..."):
                prompt_bt3 = f"""
                ĐÓNG VAI TRÒ: Trợ lý thiết kế nội dung slide dạy học THCS–THPT Việt Nam chuyên nghiệp.
                NHIỆM VỤ: Chuyển toàn bộ Kế hoạch bài dạy dưới đây thành kịch bản trình chiếu PowerPoint gồm {so_slide_target} slide.
                GIÁO VIÊN: Lê Minh Tuấn - GV
                PHONG CÁCH: {kieu_thiet_ke}

                NỘI DUNG KẾ HOẠCH BÀI DẠY ĐƯỢC CUNG CẤP:
                \"\"\"
                {final_khbd_content}
                \"\"\"

                BẮT BUỘC TUÂN THỦ NGHIÊM NGẶT QUY CÁCH SLIDE:
                1. Mỗi slide bắt đầu bằng: "SLIDE: [Tên tiêu đề slide ngắn gọn]"
                2. Nội dung hiển thị trên mặt slide theo cấu trúc phân cấp:
                   - TO: [Câu hỏi, thông điệp hoặc từ khóa trọng tâm - viết hoa/chữ đậm]
                   - VỪA: [3-4 ý chính, nhiệm vụ hoặc dữ kiện thực hiện]
                   - NHỎ: [Gợi ý hình ảnh tư liệu, nguồn hoặc hướng dẫn phụ]
                3. Dưới mỗi slide BẮT BUỘC có mục: "LỜI DẪN GV: [2-3 câu ngắn gọn giáo viên nói khi giảng slide này, không hiển thị trên mặt slide]"
                4. Tách biệt hoàn toàn phần chữ cho học sinh và lời dẫn cho giáo viên. Không đưa đáp án bài tập trực tiếp lên slide giao việc.
                """
                st.session_state["res_slide_bt3"] = call_gemini(prompt_bt3, api_key)

    if st.session_state.get("res_slide_bt3"):
        st.markdown(st.session_state["res_slide_bt3"])
        pptx_bt3_bytes = export_advanced_pptx(ten_bai, st.session_state["res_slide_bt3"], "Lê Minh Tuấn - GV")
        st.download_button(
            "📥 Tải về Tệp PowerPoint Slide Bài giảng (.pptx)", 
            data=pptx_bt3_bytes, 
            file_name=f"Slide_{ten_bai}_LeMinhTuan.pptx", 
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )
